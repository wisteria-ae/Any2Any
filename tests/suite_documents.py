import os, sys
sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))) if os.path.basename(os.path.dirname(os.path.abspath(__file__))) in ("network", "packaged") else os.path.dirname(os.path.abspath(__file__)))
from _common import ROOT, WORK, synthetic_clip
import os, sys, shutil, re, json, subprocess, traceback
os.environ["QT_QPA_PLATFORM"] = "offscreen"
sys.path.insert(0, str(ROOT))
try: sys.stdout.reconfigure(encoding="utf-8")
except Exception: pass
import Any2Any_desktop as m
import docx as pydocx
from PIL import Image

W = os.path.join(str(WORK), "docs_test")
shutil.rmtree(W, ignore_errors=True); os.makedirs(W)
fails = []
def check(name, cond, detail=""):
    print(("PASS " if cond else "FAIL ") + name + (f"   [{str(detail)[:300]}]" if (detail and not cond) else ""))
    if not cond: fails.append(name)
def conv(cls, src, target, opts=None, out=None):
    out = out or os.path.join(W, "out_" + target); os.makedirs(out, exist_ok=True)
    job = m.Job(source_path=src, target_ext=target, options=opts or {}, output_dir=out)
    msgs = []
    return cls().convert(job, lambda p, s: msgs.append(s)), msgs, out

# ---------- fixture: a rich 3-page .docx ----------
img = os.path.join(W, "pic.png"); Image.new("RGB", (400, 200), (30, 120, 220)).save(img)
d = pydocx.Document()
d.add_heading("Quarterly Report", level=1)
p = d.add_paragraph("Plain text, then "); r = p.add_run("bold words"); r.bold = True; p.add_run(" and "); r2 = p.add_run("italic words"); r2.italic = True; p.add_run(".")
d.add_heading("Accents and Cyrillic", level=2)
d.add_paragraph("Café naïve résumé — Привет мир")
d.add_paragraph("first bullet", style="List Bullet"); d.add_paragraph("second bullet", style="List Bullet")
d.add_paragraph("step one", style="List Number"); d.add_paragraph("step two", style="List Number")
t = d.add_table(rows=3, cols=3); t.style = "Table Grid"
for i, row in enumerate([("Region", "Q1", "Q2"), ("North", "10", "12"), ("South", "7", "9")]):
    for j, v in enumerate(row): t.cell(i, j).text = v
d.add_picture(img)
d.add_page_break(); d.add_heading("Second Page", level=1); d.add_paragraph("PAGE-TWO-MARKER content here.")
d.add_page_break(); d.add_heading("Third Page", level=1); d.add_paragraph("PAGE-THREE-MARKER content here.")
DOCX = os.path.join(W, "report.docx"); d.save(DOCX)

import unicodedata
def pdf_text(path):
    with m.fitz.open(path) as doc:
        return unicodedata.normalize("NFKC", "\n".join(pg.get_text() for pg in doc)), doc.page_count, sum(len(pg.get_images()) for pg in doc)

# ---------- registry: new pairs present, removed pairs gone ----------
for a, b in [("pdf","docx"),("pdf","md"),("docx","pdf"),("docx","md"),("wav","m4a"),("wav","opus"),("mp4","m4a")]:
    check(f"find_conv {a}->{b}", m.find_conv(a, b) is not None)
for a, b in [("mp4","frames"),("frames","video-from-frames"),("mp4","mp4+subs"),("url","chapters-csv")]:
    check(f"removed: {a}->{b} gone", m.find_conv(a, b) is None)
for fmt in ("frames","video-from-frames","mp4+subs","chapters-csv"):
    check(f"removed format not offered: {fmt}", fmt not in m.SUPPORTED_OUTPUTS and not any(fmt in v for v in m.OUTPUT_CATEGORIES.values()))
check("every category item is a supported output", all(f in m.SUPPORTED_OUTPUTS for v in m.OUTPUT_CATEGORIES.values() for f in v))
check("every supported output is in a category", all(any(f in v for v in m.OUTPUT_CATEGORIES.values()) for f in m.SUPPORTED_OUTPUTS))
check("every supported output has a converter from something",
      all(any(f in c.produces for c in m.REGISTRY) for f in m.SUPPORTED_OUTPUTS), [f for f in m.SUPPORTED_OUTPUTS if not any(f in c.produces for c in m.REGISTRY)])

# ---------- Word -> PDF: real Word, then built-in ----------
has_word = bool(m._find_word())
print("Word installed:", has_word)
if has_word:
    out, msgs, _ = conv(m.WordToPDF, DOCX, "pdf", out=os.path.join(W, "pdf_word"))
    txt, pages, imgs = pdf_text(out)
    check("Word->PDF (Word engine) used Word", any("Microsoft Word" in s for s in msgs), msgs)
    check("Word->PDF (Word): 3 pages", pages == 3, pages)
    check("Word->PDF (Word): heading+table+cyrillic text", all(s in txt for s in ("Quarterly Report", "North", "Привет")), txt[:200])
    check("Word->PDF (Word): image kept", imgs >= 1, imgs)
    WORD_PDF = out
orig_find = m._find_word
m._find_word = lambda: None
out, msgs, _ = conv(m.WordToPDF, DOCX, "pdf", out=os.path.join(W, "pdf_builtin"))
m._find_word = orig_find
txt, pages, imgs = pdf_text(out)
check("Word->PDF (built-in) used built-in", any("built-in" in s for s in msgs), msgs)
check("Word->PDF (built-in): multi-page", pages >= 3, pages)
check("Word->PDF (built-in): headings+table+bullets+accents", all(s in txt for s in ("Quarterly Report", "Region", "South", "first bullet", "Café", "PAGE-TWO-MARKER")), txt[:300])
check("Word->PDF (built-in): Cyrillic renders", "Привет" in txt, txt[:300])
check("Word->PDF (built-in): image kept", imgs >= 1, imgs)
BUILTIN_PDF = out

# ---------- Word -> Markdown ----------
out, msgs, od = conv(m.WordToMarkdown, DOCX, "md")
md = open(out, encoding="utf-8").read()
print("---- Word->MD output ----\n" + md + "-------------------------")
check("Word->MD: ATX headings", "# Quarterly Report" in md and "## Accents and Cyrillic" in md)
check("Word->MD: bold/italic", "**bold words**" in md and ("*italic words*" in md or "_italic words_" in md))
check("Word->MD: bullets", re.search(r"^- first bullet", md, re.M) is not None)
check("Word->MD: numbered list", re.search(r"^1\. step one", md, re.M) is not None)
check("Word->MD: table as pipes", "| Region" in md and re.search(r"\|\s*-+", md) is not None)
check("Word->MD: unicode preserved", "Привет мир" in md and "Café" in md)
imgm = re.search(r"!\[[^\]]*\]\(([^)]+)\)", md)
check("Word->MD: image linked + file exists + relative", bool(imgm) and os.path.exists(os.path.join(od, imgm.group(1))) and not os.path.isabs(imgm.group(1)), imgm and imgm.group(1))

# ---------- PDF -> Word ----------
SRC_PDF = WORD_PDF if has_word else BUILTIN_PDF
out, msgs, od = conv(m.PDFToWord, SRC_PDF, "docx")
dd = pydocx.Document(out)
alltext = "\n".join(p.text for p in dd.paragraphs) + "\n" + "\n".join(c.text for t in dd.tables for r in t.rows for c in r.cells)
check("PDF->Word: opens in python-docx", True)
check("PDF->Word: body text", "Quarterly Report" in alltext and "PAGE-TWO-MARKER" in alltext, alltext[:200])
check("PDF->Word: table rebuilt", len(dd.tables) >= 1 and "South" in alltext, len(dd.tables))
check("PDF->Word: image kept", len(dd.inline_shapes) >= 1 or any("graphic" in p._p.xml for p in dd.paragraphs), len(dd.inline_shapes))
# page range
out2, _, _ = conv(m.PDFToWord, SRC_PDF, "docx", {"pages": "2"}, out=os.path.join(W, "docx_p2"))
t2 = "\n".join(p.text for p in pydocx.Document(out2).paragraphs)
check("PDF->Word: pages=2 only has page 2", "PAGE-TWO-MARKER" in t2 and "PAGE-THREE-MARKER" not in t2 and "Quarterly" not in t2, t2[:200])

# ---------- PDF -> Markdown ----------
out, msgs, od = conv(m.PDFToMarkdown, SRC_PDF, "md", out=os.path.join(W, "md_plain"))
md2 =open(out, encoding="utf-8").read()
print("---- PDF->MD output (first 900 chars) ----\n" + md2[:900] + "\n-------------------------")
check("PDF->MD: heading detected", re.search(r"^#{1,3} .*Quarterly Report", md2, re.M) is not None, md2[:150])
check("PDF->MD: text + other pages", "PAGE-TWO-MARKER" in md2 and "PAGE-THREE-MARKER" in md2)
check("PDF->MD: table as pipes", "| Region" in md2 or "|Region" in md2, md2[:400])
check("PDF->MD: no image folder by default", not any(n.endswith("_images") for n in os.listdir(od)), os.listdir(od))
out3, _, od3 = conv(m.PDFToMarkdown, SRC_PDF, "md", {"md_images": "1", "pages": "1"}, out=os.path.join(W, "md_img"))
md3 = open(out3, encoding="utf-8").read(); imgm = re.search(r"!\[[^\]]*\]\(([^)]+)\)", md3)
check("PDF->MD: images extracted with relative links", bool(imgm) and not os.path.isabs(imgm.group(1)) and os.path.exists(os.path.join(od3, imgm.group(1))), imgm and imgm.group(1))
out4, _, _ = conv(m.PDFToMarkdown, SRC_PDF, "md", {"pages": "3"}, out=os.path.join(W, "md_p3"))
md4 = open(out4, encoding="utf-8").read()
check("PDF->MD: pages=3 only page 3", "PAGE-THREE-MARKER" in md4 and "PAGE-TWO-MARKER" not in md4)

# ---------- error handling ----------
def raises(fn, needle):
    try: fn(); return False, "no exception"
    except m.ConvertError as e: return needle.lower() in str(e).lower(), str(e)
    except Exception as e: return False, f"{type(e).__name__}: {e}"
ok, msg = raises(lambda: conv(m.PDFToWord, SRC_PDF, "docx", {"pages": "9"}, out=os.path.join(W, "e1")), "outside this PDF"); check("error: page out of range", ok, msg)
ok, msg = raises(lambda: conv(m.PDFToMarkdown, SRC_PDF, "md", {"pages": "abc"}, out=os.path.join(W, "e2")), "Bad page range"); check("error: bad page spec", ok, msg)
locked = os.path.join(W, "locked.pdf")
with m.fitz.open(SRC_PDF) as dd0: dd0.save(locked, encryption=m.fitz.PDF_ENCRYPT_AES_256, owner_pw="x", user_pw="y")
ok, msg = raises(lambda: conv(m.PDFToWord, locked, "docx", out=os.path.join(W, "e3")), "password"); check("error: password-protected PDF", ok, msg)
scan = os.path.join(W, "scan.png")
im = Image.new("RGB", (600, 400), (255, 255, 255)); im.save(scan)    # source image for an image-only PDF
with m.fitz.open() as sd:
    pg = sd.new_page(width=600, height=400); pg.insert_image(pg.rect, filename=scan); sd.save(os.path.join(W, "scanned.pdf"))
ok, msg = raises(lambda: conv(m.PDFToMarkdown, os.path.join(W, "scanned.pdf"), "md", out=os.path.join(W, "e4")), "no text was found"); check("error: a blank scanned page (nothing for OCR to read) is explained", ok, msg)
fake = os.path.join(W, "fake.docx"); open(fake, "w").write("not a zip")
ok, msg = raises(lambda: conv(m.WordToMarkdown, fake, "md", out=os.path.join(W, "e5")), "real .docx"); check("error: corrupt docx (md)", ok, msg)
m._find_word = lambda: None
ok, msg = raises(lambda: conv(m.WordToPDF, fake, "pdf", out=os.path.join(W, "e6")), "failed"); check("error: corrupt docx (pdf)", ok, msg)
m._find_word = orig_find
job = m.Job(source_path=SRC_PDF, target_ext="docx", options={}, output_dir=os.path.join(W, "e7")); os.makedirs(job.output_dir, exist_ok=True); job.cancel_flag = True
ok, msg = raises(lambda: m.PDFToWord().convert(job, lambda p, s: None), "Cancelled"); check("cancel: PDF->Word", ok, msg)
job.cancel_flag = True
ok, msg = raises(lambda: m.PDFToMarkdown().convert(job, lambda p, s: None), "Cancelled"); check("cancel: PDF->Markdown", ok, msg)
ok, msg = raises(lambda: m.WordToPDF().convert(m.Job(source_path=DOCX, target_ext="pdf", options={}, output_dir=os.path.join(W, "e8"), cancel_flag=True), lambda p, s: None), "Cancelled"); check("cancel: Word->PDF", ok, msg)

print("\n" + ("ALL PASSED" if not fails else f"{len(fails)} FAILED: {fails}"))
